"""Excel (XLSX) spreadsheet document renderer for DocuGen AI."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any, Optional
import openpyxl
from openpyxl.styles import Alignment as XlAlignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from docugen.core.document_ir import (
    ClauseIR,
    Document,
    Heading,
    ListBlock,
    Paragraph,
    Section,
    SignatureBlock,
    Table,
)
from docugen.core.models import RenderResult
from docugen.rendering.base import Renderer
from docugen.rendering.formatting import DocumentFormatting, get_default_formatting

logger = logging.getLogger("docugen.rendering.xlsx")


class XlsxRenderer(Renderer):
    """Renders Document IR to an Excel spreadsheet workbook (.xlsx)."""

    def __init__(self, formatting: Optional[DocumentFormatting] = None) -> None:
        self.formatting = formatting or get_default_formatting()

    def render(self, document: Document, destination: Path) -> RenderResult:
        """Render Document IR to XLSX at destination path."""
        try:
            dest_file = Path(destination)
            dest_file.parent.mkdir(parents=True, exist_ok=True)

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = (document.metadata.document_type or "Sheet1")[:30]

            self._populate_sheet(ws, document)

            wb.save(str(dest_file))
            file_size = dest_file.stat().st_size if dest_file.exists() else 0

            return RenderResult(
                success=True,
                output_path=str(dest_file.resolve()),
                format="xlsx",
                file_size_bytes=file_size,
            )
        except Exception as exc:
            logger.exception("Error during XLSX rendering: %s", exc)
            return RenderResult(
                success=False,
                output_path=str(destination),
                format="xlsx",
                error=str(exc),
            )

    def _populate_sheet(self, ws: Any, doc: Document) -> None:
        # Styles
        title_font = Font(name="Segoe UI", size=16, bold=True, color="1E293B")
        section_font = Font(name="Segoe UI", size=13, bold=True, color="0F172A")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        regular_font = Font(name="Segoe UI", size=10, color="334155")
        bold_font = Font(name="Segoe UI", size=10, bold=True, color="1E293B")

        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        current_row = 1

        # 1. Document Title
        if doc.metadata.title:
            ws.cell(row=current_row, column=1, value=doc.metadata.title).font = title_font
            current_row += 2

        # 2. Watermark status indicator if present
        if doc.watermark:
            wm_text = (doc.watermark.status or doc.watermark.text or "").strip().upper()
            if wm_text:
                c = ws.cell(row=current_row, column=1, value=f"STATUS: {wm_text}")
                c.font = Font(name="Segoe UI", size=10, bold=True, color="94A3B8")
                current_row += 1

        # 3. Document Sections
        for section in doc.sections:
            if section.title:
                ws.cell(row=current_row, column=1, value=section.title).font = section_font
                current_row += 1

            for elem in section.elements:
                if isinstance(elem, Heading):
                    ws.cell(row=current_row, column=1, value=elem.text).font = bold_font
                    current_row += 1

                elif isinstance(elem, Paragraph):
                    ws.cell(row=current_row, column=1, value=elem.plain_text).font = regular_font
                    current_row += 1

                elif isinstance(elem, ClauseIR):
                    ws.cell(
                        row=current_row, column=1, value=f"{elem.number or ''} {elem.title}".strip()
                    ).font = bold_font
                    current_row += 1
                    for p in elem.body:
                        ws.cell(row=current_row, column=1, value=p.plain_text).font = regular_font
                        current_row += 1

                elif isinstance(elem, ListBlock):
                    for item in elem.items:
                        ws.cell(row=current_row, column=1, value=f"• {item.plain_text}").font = regular_font
                        current_row += 1

                elif isinstance(elem, Table):
                    current_row += 1
                    for r_idx, row in enumerate(elem.rows):
                        for c_idx, cell in enumerate(row.cells, start=1):
                            xl_cell = ws.cell(
                                row=current_row, column=c_idx, value=self._parse_cell_value(cell.plain_text)
                            )
                            xl_cell.border = thin_border
                            if row.is_header or cell.bold:
                                xl_cell.fill = header_fill
                                xl_cell.font = header_font
                            else:
                                xl_cell.font = regular_font

                            # Check for right alignment or numbers
                            if isinstance(xl_cell.value, (int, float)):
                                xl_cell.alignment = XlAlignment(horizontal="right")
                        current_row += 1
                    current_row += 1

                elif isinstance(elem, SignatureBlock):
                    current_row += 1
                    if elem.intro_text:
                        ws.cell(row=current_row, column=1, value=elem.intro_text).font = regular_font
                        current_row += 1
                    for s_idx, signer in enumerate(elem.signers, start=1):
                        sig_cell = ws.cell(row=current_row, column=(s_idx * 2) - 1, value=f"Signer: {signer.name}")
                        sig_cell.font = bold_font
                        if signer.title:
                            ws.cell(row=current_row + 1, column=(s_idx * 2) - 1, value=signer.title).font = regular_font
                    current_row += 3

            current_row += 1

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if "\n" not in val and len(val) < 60:
                    max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    @staticmethod
    def _parse_cell_value(text: str) -> Any:
        clean = text.strip()
        # Clean currency
        if clean.startswith("$") or clean.startswith("€") or clean.startswith("£"):
            num_part = clean[1:].replace(",", "").strip()
            try:
                return float(num_part) if "." in num_part else int(num_part)
            except ValueError:
                pass
        try:
            if "." in clean:
                return float(clean.replace(",", ""))
            return int(clean.replace(",", ""))
        except ValueError:
            return text
