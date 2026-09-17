"""Parser and merger for filesystem DOCX templates with placeholder variables."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict
import docx

from docugen.core.document_ir import (
    Alignment,
    Document,
    DocumentMetadata,
    Heading,
    Paragraph,
    Section,
    Table,
    TableCell,
    TableRow,
    TextRun,
)
from docugen.templates.loader import get_safe_loader
from docugen.templates.registry import TemplateDefinition


def parse_and_merge_docx_template(file_path: Path, data: Dict[str, Any], template_def: TemplateDefinition) -> Document:
    """Parse a DOCX template, interpolate Jinja2 expressions with data, and return Document IR."""
    doc_file = docx.Document(str(file_path))
    safe_loader = get_safe_loader()

    doc_ir = Document(
        metadata=DocumentMetadata(
            title=str(data.get("title", template_def.name.replace("_", " ").title())),
            document_type=template_def.document_type,
            template_name=template_def.name,
        )
    )

    section_elements = []

    for p in doc_file.paragraphs:
        if not p.text.strip():
            continue

        merged_text = safe_loader.render_string(p.text, data)
        # Check if heading
        style_name = p.style.name.lower()
        if "heading 1" in style_name:
            section_elements.append(Heading(text=merged_text, level=1))
        elif "heading 2" in style_name:
            section_elements.append(Heading(text=merged_text, level=2))
        elif "heading 3" in style_name:
            section_elements.append(Heading(text=merged_text, level=3))
        else:
            # Detect formatting from runs
            runs = []
            for r in p.runs:
                if r.text:
                    merged_run_text = safe_loader.render_string(r.text, data)
                    runs.append(
                        TextRun(
                            text=merged_run_text,
                            bold=bool(r.bold),
                            italic=bool(r.italic),
                            underline=bool(r.underline),
                        )
                    )
            if not runs:
                runs = [TextRun(text=merged_text)]
            section_elements.append(Paragraph(runs=runs))

    for t in doc_file.tables:
        table_rows = []
        for row in t.rows:
            row_cells = []
            for cell in row.cells:
                merged_cell = safe_loader.render_string(cell.text.strip(), data)
                row_cells.append(TableCell(content=merged_cell))
            table_rows.append(TableRow(cells=row_cells))
        if table_rows:
            section_elements.append(Table(rows=table_rows, bordered=True))

    doc_ir.sections.append(
        Section(
            title=template_def.name.replace("_", " ").title(),
            elements=section_elements,
        )
    )
    return doc_ir
