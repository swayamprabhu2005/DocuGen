"""Tests for PDF and DOCX rendering."""

import pytest
from pathlib import Path

from docugen.core.document_ir import (
    Alignment,
    ClauseIR,
    Document,
    DocumentMetadata,
    Footer,
    Header,
    Heading,
    ListBlock,
    ListItem,
    Paragraph,
    Section,
    SignatureBlock,
    Signer,
    Table,
    TableCell,
    TableRow,
    Watermark,
)
from docugen.rendering.docx import DocxRenderer
from docugen.rendering.pdf import PdfRenderer


def _make_rich_document() -> Document:
    """Build a complex Document IR for rendering tests."""
    return Document(
        metadata=DocumentMetadata(
            title="Rendering Test Document",
            author="DocuGen Tests",
            document_type="test",
        ),
        header=Header(left_text="DocuGen Test", right_text="DRAFT"),
        footer=Footer(include_page_number=True, left_text="Confidential"),
        watermark=Watermark(text="SAMPLE"),
        sections=[
            Section(
                title="Section 1: Overview",
                elements=[
                    Heading(text="1.1 Introduction", level=2),
                    Paragraph.from_text(
                        "This document is generated for rendering validation purposes.",
                        space_after=8.0,
                    ),
                    Paragraph.from_text("Bold text and italic text.", bold=False, space_after=6.0),
                    ListBlock(
                        ordered=False,
                        items=[
                            ListItem.from_text("Bullet point one", prefix="•"),
                            ListItem.from_text("Bullet point two", prefix="•"),
                            ListItem.from_text("Bullet point three", prefix="•"),
                        ],
                    ),
                ],
            ),
            Section(
                title="Section 2: Data Tables",
                elements=[
                    Table(
                        has_header=True,
                        bordered=True,
                        col_widths=[0.4, 0.3, 0.3],
                        rows=[
                            TableRow(
                                is_header=True,
                                cells=[
                                    TableCell(content="Description", bold=True),
                                    TableCell(content="Quantity", bold=True, alignment=Alignment.RIGHT),
                                    TableCell(content="Price", bold=True, alignment=Alignment.RIGHT),
                                ],
                            ),
                            TableRow(
                                cells=[
                                    TableCell(content="Widget Alpha"),
                                    TableCell(content="5", alignment=Alignment.RIGHT),
                                    TableCell(content="$250.00", alignment=Alignment.RIGHT),
                                ]
                            ),
                            TableRow(
                                cells=[
                                    TableCell(content="Widget Beta"),
                                    TableCell(content="3", alignment=Alignment.RIGHT),
                                    TableCell(content="$150.00", alignment=Alignment.RIGHT),
                                ]
                            ),
                        ],
                    ),
                ],
            ),
            Section(
                title="Section 3: Signature Block",
                elements=[
                    ClauseIR(
                        clause_id="test_clause",
                        title="1.0 Test Clause",
                        body=[
                            Paragraph.from_text("This clause confirms the rendering of a ClauseIR element."),
                        ],
                    ),
                    SignatureBlock(
                        signers=[
                            Signer(name="Alice Anderson", title="CEO", organization="Alpha Corp", date="2026-09-18"),
                            Signer(name="Bob Baker", title="CTO", organization="Beta LLC", date="2026-09-18"),
                        ],
                        layout="side_by_side",
                        intro_text="IN WITNESS WHEREOF, the parties hereto have executed this Agreement.",
                    ),
                ],
            ),
        ],
    )


class TestPdfRenderer:

    def test_pdf_renders_successfully(self, temp_output_dir):
        doc = _make_rich_document()
        renderer = PdfRenderer()
        dest = temp_output_dir / "test_render.pdf"
        result = renderer.render(doc, dest)

        assert result.success, f"PDF render failed: {result.error}"
        assert dest.exists()
        assert dest.stat().st_size > 2000

    def test_pdf_with_watermark(self, temp_output_dir):
        doc = _make_rich_document()
        doc.watermark = Watermark(text="CONFIDENTIAL", opacity=0.3)
        renderer = PdfRenderer()
        dest = temp_output_dir / "test_watermark.pdf"
        result = renderer.render(doc, dest)
        assert result.success
        assert dest.exists()

    def test_pdf_without_header_footer(self, temp_output_dir):
        doc = Document(
            metadata=DocumentMetadata(title="Minimal PDF"),
            sections=[
                Section(elements=[Paragraph.from_text("Simple minimal paragraph.")])
            ],
        )
        renderer = PdfRenderer()
        dest = temp_output_dir / "minimal.pdf"
        result = renderer.render(doc, dest)
        assert result.success
        assert dest.exists()

    def test_pdf_result_metadata(self, temp_output_dir):
        doc = _make_rich_document()
        renderer = PdfRenderer()
        dest = temp_output_dir / "meta_test.pdf"
        result = renderer.render(doc, dest)
        assert result.format == "pdf"
        assert result.file_size_bytes > 0


class TestDocxRenderer:

    def test_docx_renders_successfully(self, temp_output_dir):
        doc = _make_rich_document()
        renderer = DocxRenderer()
        dest = temp_output_dir / "test_render.docx"
        result = renderer.render(doc, dest)

        assert result.success, f"DOCX render failed: {result.error}"
        assert dest.exists()
        assert dest.stat().st_size > 5000

    def test_docx_contains_valid_file(self, temp_output_dir):
        doc = _make_rich_document()
        renderer = DocxRenderer()
        dest = temp_output_dir / "valid_docx.docx"
        renderer.render(doc, dest)

        # Verify it's a valid docx by loading it
        import docx as python_docx
        loaded = python_docx.Document(str(dest))
        assert len(loaded.paragraphs) > 0

    def test_docx_result_metadata(self, temp_output_dir):
        doc = _make_rich_document()
        renderer = DocxRenderer()
        dest = temp_output_dir / "meta_test.docx"
        result = renderer.render(doc, dest)
        assert result.format == "docx"
        assert result.file_size_bytes > 0
